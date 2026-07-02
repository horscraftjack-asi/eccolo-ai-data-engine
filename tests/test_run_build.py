"""Full run_build() integration test with ANTHROPIC_API_KEY unset.

Proves the deterministic flow end-to-end: parse config -> read Meta + YouTube CSVs ->
validate -> score -> build the .xlsx with labelled insight shells (no Claude), and that the
graceful-degradation note is exactly what the app surfaces to the result screen.
"""
import os

import pytest
from openpyxl import load_workbook

from core import build_workbook as bw

FIXTURES = os.path.join(os.path.dirname(__file__), "fixtures")
CONFIG = os.path.join(FIXTURES, "testclient-analytics-config.md")

# The exact string core/insights.py returns when the key is absent. app.py greps notes for
# "insight"/"claude" to populate X-Insight-Note, so this wording is load-bearing.
NO_KEY_NOTE = "No Claude insights (ANTHROPIC_API_KEY not set)."


@pytest.fixture
def no_api_key(monkeypatch):
    monkeypatch.delenv("ANTHROPIC_API_KEY", raising=False)


def test_run_build_meta_and_youtube_no_key(no_api_key, tmp_path):
    result = bw.run_build(
        config_path=CONFIG,
        csv_paths={
            "facebook": os.path.join(FIXTURES, "facebook.csv"),
            "instagram": os.path.join(FIXTURES, "instagram.csv"),
            "youtube": os.path.join(FIXTURES, "youtube.csv"),
        },
        out_dir=str(tmp_path),
    )

    assert result["status"] == "built"
    assert result["month"] == "May 2026"

    # Workbook was produced and is a real, openable .xlsx
    out = result["output"]
    assert out.endswith("TestClient_May2026_Performance.xlsx")
    assert os.path.exists(out)
    wb = load_workbook(out)
    # data tabs + insight shell tabs all present
    assert len(wb.sheetnames) >= 4

    # YouTube combined export auto-split by Duration: 2 shorts (<=60s), 2 long-form (>60s)
    assert result["counts"]["youtube_shorts"] == 2
    assert result["counts"]["youtube_longform"] == 2
    assert result["counts"]["instagram"] == 5
    assert result["counts"]["facebook"] == 4

    # Graceful degradation: exact "No Claude insights" note, shells (not Claude text) in the tabs
    assert NO_KEY_NOTE in result["notes"]

    # Cross-platform leaderboard is populated from the scored tables
    assert result["top_posts"]
    assert all("score" in row for row in result["top_posts"])


def test_run_build_emits_contract_provenance(no_api_key, tmp_path):
    """Phase 1: run_build() stamps a contract §1.3 provenance block onto the report.
    source_ids stays [] until Phase 2 derives a source_id per scored post."""
    result = bw.run_build(
        config_path=CONFIG,
        csv_paths={"instagram": os.path.join(FIXTURES, "instagram.csv")},
        out_dir=str(tmp_path),
    )
    prov = result["provenance"]
    assert set(prov) == {"run_id", "generated_at", "tool", "tool_version",
                         "client_slug", "source_ids"}
    assert prov["tool"] == "analytics"
    assert prov["tool_version"] == "1.0.0"
    assert prov["client_slug"] == "testclient"
    assert prov["source_ids"] == []
    assert prov["run_id"].startswith("analytics-")


def test_run_build_validate_only_no_key(no_api_key, tmp_path):
    """validate_only short-circuits before scoring/insights and reports row counts."""
    result = bw.run_build(
        config_path=CONFIG,
        csv_paths={"instagram": os.path.join(FIXTURES, "instagram.csv")},
        out_dir=str(tmp_path),
        validate_only=True,
    )
    assert result["status"] == "validated"
    assert result["validation"]["instagram"]["rows"] == 5
    assert result["validation"]["instagram"]["missing_columns"] == []
